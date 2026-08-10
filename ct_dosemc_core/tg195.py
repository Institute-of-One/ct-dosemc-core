"""AAPM TG-195 Case 4 -- code verification against a published reference.

Why this is the Phase-1 gate
----------------------------
A measured scanner CTDIvol cannot verify this code: the simulation has no bowtie,
a flat fan and a homogeneous PMMA cylinder, so any disagreement is unattributable
between a coding error and a missing physical component.  TG-195 Case 4 fixes
that by *specifying* exactly the configuration we already have -- no bowtie, a
PMMA cylinder of radius 160 mm, an isotropic point source at 600 mm, and energy
scored directly in PMMA in **eV per emitted photon**.  Four independent Monte
Carlo packages (EGSnrc, Geant4, MCNP, PENELOPE) published results for it.

Three problems disappear at this gate, none of them by being ignored:

* **no bowtie is needed** -- TG-195 Case 4 does not have one;
* **no PMMA-to-air-kerma conversion is needed** -- energy deposited in PMMA *is*
  the scored quantity;
* **no absolute mAs normalisation is needed** -- the unit is per emitted photon,
  so SpekPy's absolute tube output drops out of the comparison entirely.

Licensing
---------
The specification *values* below are transcribed from the report's prose because
they are a specification -- the dimensions of a phantom anyone may build.  The
report itself, its distributed spectra and its tabulated reference results are
copyrighted (© 2015 AAPM) and are **not** committed: they are fetched by
``tools/fetch_tg195.py`` and read from ``data/tg195/`` at run time.

Reference: Sechopoulos et al., Med. Phys. 42(10), 5679-5691 (2015),
report §3.4 (pp. 25-27) and §4.4.
"""

from __future__ import annotations

import math
from dataclasses import dataclass
from pathlib import Path

import numpy as np

from .phantom import Material, VoxelPhantom, cylinder_occupancy

__all__ = [
    "CASE4",
    "Case4Spec",
    "SubFan",
    "build_case4_phantom",
    "first_test_voi_masks",
    "first_test_voi_weights",
    "perimeter_replica_weights",
    "read_case4_reference",
    "read_tg195_spectrum",
    "second_test_voi_masks",
    "second_test_voi_weights",
    "sub_fans",
    "write_monoenergetic_spectrum",
]


@dataclass(frozen=True)
class Case4Spec:
    """The Case 4 geometry, exactly as §3.4 states it. Lengths in cm.

    ``phantom_length_cm`` defaults to the specified 300 cm ("to approximate an
    infinite cylinder").  It is a parameter because that length costs 22 million
    voxels at 2.5 mm, and whether the ends matter is a convergence question with
    an answer, not an article of faith -- ``tools/run_tg195.py --length-scan``
    measures it.
    """

    phantom_radius_cm: float = 16.0
    phantom_length_cm: float = 300.0
    sad_cm: float = 60.0
    fan_width_at_isocentre_cm: float = 32.0
    """Field width in the fan direction, measured at the isocentre. Equal to the
    body diameter, so the fan slightly *clips* the cylinder: a ray tangent to the
    cylinder needs asin(16/60) = 15.466 deg, the field only reaches
    atan(16/60) = 14.931 deg. That is what the report specifies."""

    fan_thickness_at_isocentre_cm: float = 1.0
    """``t`` in Table 6. The report gives two values, 10 mm and 80 mm."""

    pmma_density_g_cm3: float = 1.19
    air_density_g_cm3: float = 1.20479e-3

    # -- first test: four contiguous axial sections of the whole cylinder
    first_test_segment_thickness_cm: float = 1.0
    first_test_n_segments: int = 4

    # -- second test: two 10 mm diameter PMMA cylinders, z from -50 to +50 mm
    second_test_voi_diameter_cm: float = 1.0
    second_test_voi_length_cm: float = 10.0
    second_test_perimeter_radius_cm: float = 15.0

    monoenergetic_kev: float = 56.4
    """"equivalent to the mean energy of the spectrum" (§3.4)."""

    @property
    def half_fan_deg(self) -> float:
        """atan, not asin: the field is a rectangle on the isocentre plane."""
        return math.degrees(math.atan(0.5 * self.fan_width_at_isocentre_cm / self.sad_cm))

    @property
    def fan_aperture_deg(self) -> float:
        return 2.0 * self.half_fan_deg


CASE4 = Case4Spec()

# Table 6 gives two fan thicknesses and three projection schemes.
CASE4_FAN_THICKNESSES_CM = (1.0, 8.0)
CASE4_DISCRETE_PROJECTION_ANGLES_DEG = tuple(float(a) for a in range(0, 360, 10))


# ---------------------------------------------------------------------------
# geometry
# ---------------------------------------------------------------------------


def build_case4_phantom(
    spec: Case4Spec,
    *,
    air: Material,
    pmma: Material,
    voxel_size_cm: float = 0.25,
    air_margin_cm: float = 1.0,
    air_margin_z_cm: float | None = None,
    length_cm: float | None = None,
) -> VoxelPhantom:
    """Voxelise the Case 4 body phantom: a PMMA cylinder in air.

    ``length_cm`` overrides ``spec.phantom_length_cm`` for convergence studies.

    ``air_margin_cm`` is not cosmetic.  MC-GPU jumps each photon from the focal
    spot to the voxel bounding box in a straight line **without interaction**, so
    whatever air lies outside the box is simply not there.  TG-195 fills the rest
    of the geometry with air, and the source sits 44 cm beyond the cylinder
    surface, so a tight box silently omits 44 cm of attenuation and of in-air
    scatter.  Measured effect on Case 4: about 0.8% on the scored segments.

    The transverse margin is what matters -- the beam travels in the xy-plane --
    so ``air_margin_z_cm`` is separate and can stay small, which is the
    difference between an 9-million-voxel geometry and a 130-million-voxel one.
    """
    v = voxel_size_cm
    length = spec.phantom_length_cm if length_cm is None else length_cm
    margin_z = air_margin_cm if air_margin_z_cm is None else air_margin_z_cm
    span_xy = 2.0 * spec.phantom_radius_cm + 2.0 * air_margin_cm
    span_z = length + 2.0 * margin_z
    nx = ny = int(math.ceil(span_xy / v))
    nz = int(math.ceil(span_z / v))

    x = (np.arange(nx) + 0.5) * v
    y = (np.arange(ny) + 0.5) * v
    z = (np.arange(nz) + 0.5) * v
    cx, cy, cz = 0.5 * nx * v, 0.5 * ny * v, 0.5 * nz * v

    in_disc = ((x[None, :] - cx) ** 2 + (y[:, None] - cy) ** 2) <= spec.phantom_radius_cm**2
    in_z = np.abs(z - cz) <= 0.5 * length
    inside = in_z[:, None, None] & in_disc[None, :, :]

    return VoxelPhantom(
        material=np.where(inside, pmma.index, air.index).astype(np.int16),
        density=np.where(inside, spec.pmma_density_g_cm3, spec.air_density_g_cm3),
        voxel_size_cm=(v, v, v),
        materials=(air, pmma),
    )


def _crop(full: np.ndarray, origin_zyx, shape_zyx) -> np.ndarray:
    z0, y0, x0 = origin_zyx
    dz, dy, dx = shape_zyx
    return full[z0 : z0 + dz, y0 : y0 + dy, x0 : x0 + dx]


def first_test_voi_weights(
    phantom: VoxelPhantom,
    spec: Case4Spec,
    *,
    roi_origin_zyx: tuple[int, int, int] = (0, 0, 0),
    roi_shape_zyx: tuple[int, int, int] | None = None,
    subdivisions: int = 64,
) -> dict[str, np.ndarray]:
    """The four contiguous axial sections of the first test, as occupancy arrays.

    "Each cylindrical segment is an axial tomographic section of the phantom with
    thickness 10 mm, with centers located at z = 0 mm, z = 10 mm, z = 20 mm, and
    z = 30 mm" (§3.4).  Each is the *whole* disc, so the scoring mass is large and
    the statistics are cheap.

    Values are the fraction of each voxel inside the segment, in ``[0, 1]``.

    **The radial boundary is deliberately NOT partial-volume weighted here**, and
    that is not an oversight.  This VOI's edge *is* the phantom's edge, and the
    phantom was voxelised by the same voxel-centre test, so a voxel whose centre
    is inside the cylinder is entirely PMMA in the simulated geometry and its full
    mass genuinely absorbed dose.  Down-weighting it would discard real deposited
    energy -- and precisely at the entrance surface, where the dose is highest, so
    the error is much larger than the volume mismatch suggests (measured: -2.2% on
    the segment holding the primary beam).

    The residual +0.19% mismatch between the voxelised disc and the analytic one
    is a *geometry* discretisation error that affects transport as well as
    scoring.  It cannot be corrected by re-weighting the score; it converges with
    voxel size, which is what the G3 scan measures.

    The z boundaries sit inside bulk PMMA, so those *are* weighted exactly.  The
    second test is the opposite case -- a small VOI wholly inside uniform PMMA,
    with no material boundary to respect -- and is fully weighted there.
    """
    if roi_shape_zyx is None:
        roi_shape_zyx = phantom.shape_zyx
    thickness = spec.first_test_segment_thickness_cm
    _cx, _cy, cz = phantom.centre_cm()
    # The phantom's own PMMA voxels are the radial extent, by construction.
    pmma_index = max(m.index for m in phantom.materials)
    is_pmma = phantom.material == pmma_index

    nz = phantom.shape_zyx[0]
    vz = phantom.voxel_size_cm[2]
    z = (np.arange(nz) + 0.5) * vz

    out: dict[str, np.ndarray] = {}
    for i in range(spec.first_test_n_segments):
        centre = cz + i * thickness
        lo = np.maximum(z - 0.5 * vz, centre - 0.5 * thickness)
        hi = np.minimum(z + 0.5 * vz, centre + 0.5 * thickness)
        f_z = np.clip(hi - lo, 0.0, None) / vz
        w = _crop(f_z[:, None, None] * is_pmma, roi_origin_zyx, roi_shape_zyx)
        if not w.any():
            raise ValueError(f"first-test segment {i + 1} falls outside the dose tally ROI")
        out[f"segment_{i + 1}"] = w
    return out


def first_test_voi_masks(
    phantom: VoxelPhantom,
    spec: Case4Spec,
    *,
    roi_origin_zyx: tuple[int, int, int] = (0, 0, 0),
    roi_shape_zyx: tuple[int, int, int] | None = None,
) -> dict[str, np.ndarray]:
    """Boolean form of :func:`first_test_voi_weights` -- every touched voxel."""
    return {
        k: w > 0.0
        for k, w in first_test_voi_weights(
            phantom, spec, roi_origin_zyx=roi_origin_zyx, roi_shape_zyx=roi_shape_zyx
        ).items()
    }


def second_test_voi_weights(
    phantom: VoxelPhantom,
    spec: Case4Spec,
    *,
    roi_origin_zyx: tuple[int, int, int] = (0, 0, 0),
    roi_shape_zyx: tuple[int, int, int] | None = None,
    subdivisions: int = 64,
) -> dict[str, np.ndarray]:
    """The two 10 mm diameter, 100 mm long PMMA cylinders, as occupancy arrays.

    One on the axis, one with its centre 150 mm from the axis (§3.4).

    **This is where the partial volume matters.** A 10 mm diameter circle on a
    2.5 mm grid resolves, under a voxel-centre test, as exactly 12 voxels =
    0.750 cm^2 against a true 0.785 cm^2 -- 4.51% short.  Both VOI centres sit on
    voxel corners, so both lose the same 4.51%, and because the score is energy
    deposited (proportional to selected volume) it comes straight off the result.
    """
    if roi_shape_zyx is None:
        roi_shape_zyx = phantom.shape_zyx
    cx, cy, cz = phantom.centre_cm()
    half_length = 0.5 * spec.second_test_voi_length_cm
    radius = 0.5 * spec.second_test_voi_diameter_cm

    offsets = {"center": (0.0, 0.0), "perimeter": (0.0, spec.second_test_perimeter_radius_cm)}
    out: dict[str, np.ndarray] = {}
    for name, (ox, oy) in offsets.items():
        full = cylinder_occupancy(
            phantom.n_voxels,
            phantom.voxel_size_cm,
            centre_xy_cm=(cx + ox, cy + oy),
            radius_cm=radius,
            z_min_cm=cz - half_length,
            z_max_cm=cz + half_length,
            subdivisions=subdivisions,
        )
        w = _crop(full, roi_origin_zyx, roi_shape_zyx)
        if not w.any():
            raise ValueError(f"second-test VOI {name!r} falls outside the dose tally ROI")
        out[name] = w
    return out


def perimeter_replica_weights(
    phantom: VoxelPhantom,
    spec: Case4Spec,
    *,
    n_positions: int | None = None,
    roi_origin_zyx: tuple[int, int, int] = (0, 0, 0),
    roi_shape_zyx: tuple[int, int, int] | None = None,
    subdivisions: int = 64,
) -> dict[str, np.ndarray]:
    """The perimeter VOI replicated around its circle of symmetry.

    After a full 360 degree rotation the dose field in a cylindrically symmetric
    phantom depends only on ``(r, z)``, never on azimuth.  The perimeter VOI can
    therefore be evaluated at **any** azimuth on its 150 mm circle and the
    expectation is identical -- so evaluating it at many non-overlapping azimuths
    and averaging costs no extra Monte Carlo and reduces the variance by the
    number of positions.

    The circle is 2*pi*150 = 942 mm around and the VOI is 10 mm across, so 94
    positions fit without overlapping.  That is a factor of ~9.7 on the standard
    deviation, from an existing tally.

    It is also a *second, independent* uncertainty estimate: the sample SD across
    the positions must agree with the replica SD from independent seeds. If it
    does not, something breaks the symmetry that should not -- which is exactly
    the structure of the four-chamber CTDI check that caught the PRNG defect.

    The centre VOI has no such freedom: it sits on the rotation axis, so there is
    no azimuth to replicate over. Only the perimeter gets this for free.
    """
    circumference = 2.0 * math.pi * spec.second_test_perimeter_radius_cm
    max_positions = int(circumference // spec.second_test_voi_diameter_cm)
    if n_positions is None:
        n_positions = max_positions
    if n_positions > max_positions:
        raise ValueError(
            f"{n_positions} positions of {spec.second_test_voi_diameter_cm:g} cm "
            f"would overlap on a {circumference:.1f} cm circle; at most {max_positions}"
        )

    if roi_shape_zyx is None:
        roi_shape_zyx = phantom.shape_zyx
    cx, cy, cz = phantom.centre_cm()
    half_length = 0.5 * spec.second_test_voi_length_cm
    r = spec.second_test_perimeter_radius_cm

    out: dict[str, np.ndarray] = {}
    for m in range(n_positions):
        theta = 2.0 * math.pi * m / n_positions
        full = cylinder_occupancy(
            phantom.n_voxels,
            phantom.voxel_size_cm,
            centre_xy_cm=(cx + r * math.sin(theta), cy + r * math.cos(theta)),
            radius_cm=0.5 * spec.second_test_voi_diameter_cm,
            z_min_cm=cz - half_length,
            z_max_cm=cz + half_length,
            subdivisions=subdivisions,
        )
        w = _crop(full, roi_origin_zyx, roi_shape_zyx)
        if not w.any():
            raise ValueError(
                f"perimeter replica {m} (azimuth {math.degrees(theta):.1f} deg) "
                "falls outside the dose tally ROI"
            )
        out[f"perimeter_{m:03d}"] = w
    return out


def second_test_voi_masks(
    phantom: VoxelPhantom,
    spec: Case4Spec,
    *,
    roi_origin_zyx: tuple[int, int, int] = (0, 0, 0),
    roi_shape_zyx: tuple[int, int, int] | None = None,
) -> dict[str, np.ndarray]:
    """Boolean form of :func:`second_test_voi_weights` -- every touched voxel."""
    return {
        k: w > 0.0
        for k, w in second_test_voi_weights(
            phantom, spec, roi_origin_zyx=roi_origin_zyx, roi_shape_zyx=roi_shape_zyx
        ).items()
    }


# ---------------------------------------------------------------------------
# the fan: a rectangle on the isocentre plane, not a (cos theta, phi) rectangle
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class SubFan:
    """One azimuthal slice of the fan, with its own polar aperture.

    TG-195 collimates the source to a field that is a **rectangle on the
    isocentre plane** -- 320 mm wide by ``t`` thick.  MC-GPU instead samples a
    rectangle in ``(cos theta, phi)``, whose z extent at the isocentre plane
    grows as ``1/cos(psi)`` towards the fan edges.  Emitting into the larger
    window dilutes every "eV per emitted photon" figure by the solid-angle ratio

        Omega_spec / Omega_mcgpu = sin(psi_max) / psi_max

    which is 1.1% for this geometry -- the whole error budget of the test.

    Splitting the fan into K azimuthal sub-fans and giving each the polar
    aperture that reproduces thickness ``t`` at *its* azimuth removes the
    discrepancy to O(1/K^2).  The same decomposition is what a bowtie needs later
    (a per-sub-fan spectrum), so this is not scaffolding built twice.
    """

    index: int
    centre_psi_deg: float
    """Azimuth of this sub-fan's central ray, from the fan centre."""

    azimuthal_aperture_deg: float
    polar_aperture_deg: float

    @property
    def solid_angle_sr(self) -> float:
        return (
            2.0
            * math.sin(0.5 * math.radians(self.polar_aperture_deg))
            * math.radians(self.azimuthal_aperture_deg)
        )

    def direction(self) -> tuple[float, float, float]:
        """Central direction, for a source whose fan centre points along +y.

        MC-GPU samples phi about the source direction, so a sub-fan is aimed by
        rotating the source direction rather than by shifting the phi window.
        """
        psi = math.radians(self.centre_psi_deg)
        return (-math.sin(psi), math.cos(psi), 0.0)


def sub_fans(spec: Case4Spec, n: int = 1, *, thickness_cm: float | None = None) -> list[SubFan]:
    """Split the Case 4 fan into ``n`` azimuthal sub-fans.

    ``n = 1`` reproduces the naive single-aperture geometry, which is useful
    precisely because the difference against larger ``n`` is measurable.
    """
    if n < 1:
        raise ValueError("need at least one sub-fan")
    t = spec.fan_thickness_at_isocentre_cm if thickness_cm is None else thickness_cm
    psi_max = spec.half_fan_deg
    edges = np.linspace(-psi_max, psi_max, n + 1)
    fans = []
    for k in range(n):
        centre = 0.5 * float(edges[k] + edges[k + 1])
        # Thickness t at the isocentre plane, at this azimuth, subtends
        #   2 * atan( (t/2) * cos(psi) / SAD )
        # because the ray to that point travels SAD/cos(psi), not SAD.
        polar = 2.0 * math.degrees(
            math.atan(0.5 * t * math.cos(math.radians(centre)) / spec.sad_cm)
        )
        fans.append(
            SubFan(
                index=k,
                centre_psi_deg=centre,
                azimuthal_aperture_deg=float(edges[k + 1] - edges[k]),
                polar_aperture_deg=polar,
            )
        )
    return fans


def specified_solid_angle_sr(spec: Case4Spec, *, thickness_cm: float | None = None) -> float:
    """Exact solid angle of the specified rectangular field.

    ``Omega = integral over psi of (t/SAD) cos(psi) dpsi = 2 (t/SAD) sin(psi_max)``.
    """
    t = spec.fan_thickness_at_isocentre_cm if thickness_cm is None else thickness_cm
    return 2.0 * (t / spec.sad_cm) * math.sin(math.radians(spec.half_fan_deg))


# ---------------------------------------------------------------------------
# spectra
# ---------------------------------------------------------------------------


def write_monoenergetic_spectrum(
    path: str | Path, energy_kev: float, *, width_ev: float = 1.0e-3
) -> dict[str, object]:
    """Write a (numerically) monoenergetic MC-GPU spectrum file.

    MC-GPU samples uniformly between consecutive bin edges and needs a
    negative-probability terminator, so a monoenergetic source is a single bin of
    negligible width.  This removes spectrum sampling, rebinning and the 255-bin
    limit from the comparison entirely -- if the monoenergetic case disagrees,
    the fault is in transport or geometry, not in the spectrum.
    """
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    e0 = energy_kev * 1.0e3
    out.write_text(
        "\n".join(
            [
                "# Monoenergetic MC-GPU spectrum, written by ct_dosemc_core",
                f"# {energy_kev:g} keV (TG-195 Case 4 monoenergetic case)",
                "# columns: lower bin energy [eV]   relative probability",
                f"{e0:.6f}   1.0",
                f"{e0 + width_ev:.6f}   -1",
                "",
            ]
        ),
        encoding="ascii",
    )
    return {"path": str(out), "energy_kev": energy_kev, "bin_width_ev": width_ev, "n_bins": 1}


def read_tg195_spectrum(
    workbook: str | Path, *, column: str = "120 kVp"
) -> tuple[np.ndarray, np.ndarray, dict[str, str]]:
    """Read a distributed TG-195 spectrum from ``03-X-Ray Spectra Definition.xlsx``.

    Returns mid-bin energies [keV], relative fluence, and the header metadata
    (target, kVp, anode angle, filtration, mean energy, HVL) that the sheet
    states for that column.

    The file is copyrighted and is not committed; ``tools/fetch_tg195.py`` puts it
    in ``data/tg195/``.
    """
    import openpyxl

    wb = openpyxl.load_workbook(Path(workbook), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]

    header_row = next(
        i for i, r in enumerate(rows) if r and str(r[0] or "").startswith("Mid-bin Energy")
    )
    col = next(
        (
            j
            for j in range(1, len(rows[2]))
            if rows[2][j] is not None and str(rows[2][j]).strip() == column
        ),
        None,
    )
    if col is None:
        available = [str(c) for c in rows[2][1:] if c]
        raise KeyError(f"spectrum column {column!r} not found; available: {available}")

    meta_labels = ("target", "kvp", "anode_angle", "ripple", "filtration", "mean_energy", "hvl")
    meta = {
        label: str(rows[i][col]).strip()
        for label, i in zip(meta_labels, range(1, 1 + len(meta_labels)), strict=False)
        if rows[i][col] is not None
    }

    energy, fluence = [], []
    for r in rows[header_row + 1 :]:
        if not r or r[0] is None or r[col] is None:
            continue
        energy.append(float(r[0]))
        fluence.append(float(r[col]))
    if not energy:
        raise ValueError(f"no spectrum data in column {column!r}")
    return np.asarray(energy), np.asarray(fluence), meta


def write_tg195_spectrum_file(
    path: str | Path, energy_kev: np.ndarray, fluence: np.ndarray
) -> dict[str, object]:
    """Write a TG-195 distributed spectrum in MC-GPU's format.

    The sheet gives **mid-bin** energies on a uniform grid; MC-GPU wants lower
    edges plus a terminating upper edge, so the edges are reconstructed from the
    bin width rather than assumed.  Zero-fluence bins below the filtration cutoff
    are dropped, which keeps the file inside MC-GPU's ``MAX_ENERGY_BINS = 256``.
    """
    out = Path(path)
    out.parent.mkdir(parents=True, exist_ok=True)
    e = np.asarray(energy_kev, dtype=float)
    f = np.asarray(fluence, dtype=float)
    if e.size < 2:
        raise ValueError("need at least two spectrum bins")
    dk = float(np.median(np.diff(e)))

    keep = f > 0.0
    if not keep.any():
        raise ValueError("spectrum is empty")
    first, last = int(np.argmax(keep)), int(len(keep) - np.argmax(keep[::-1]) - 1)
    e, f = e[first : last + 1], f[first : last + 1]
    if e.size > 255:
        raise ValueError(
            f"{e.size} bins exceeds MC-GPU's MAX_ENERGY_BINS=256; rebin before writing"
        )

    prob = f / f.sum()
    lower = e - 0.5 * dk
    lines = [
        "# TG-195 distributed spectrum, written by ct_dosemc_core",
        "# source: AAPM Report 195 electronic resources (copyrighted, not redistributed)",
        "# columns: lower bin energy [eV]   relative probability",
    ]
    lines += [f"{lower[i] * 1e3:.4f}   {prob[i]:.8e}" for i in range(e.size)]
    lines.append(f"{(lower[-1] + dk) * 1e3:.4f}   -1")
    out.write_text("\n".join(lines) + "\n", encoding="ascii")

    return {
        "path": str(out),
        "n_bins": int(e.size),
        "bin_width_kev": dk,
        "e_min_kev": float(lower[0]),
        "e_max_kev": float(lower[-1] + dk),
        "mean_energy_kev": float(np.sum(prob * e)),
    }


# ---------------------------------------------------------------------------
# reference results
# ---------------------------------------------------------------------------

# Column of the Case 4 results sheets for each (beam, fan thickness) combination.
_CASE4_COLUMNS = {
    ("56.4 keV", 1.0): 3,
    ("56.4 keV", 8.0): 6,
    ("120 kVp", 1.0): 10,
    ("120 kVp", 8.0): 13,
}

_CODES = ("EGSnrc", "Geant4", "MCNP", "Penelope")


def _parse_range(text: str) -> tuple[float, float]:
    """'11,581.99 - 11,618.30' -> (11581.99, 11618.30)."""
    lo, _, hi = str(text).partition("-")
    return float(lo.replace(",", "").strip()), float(hi.replace(",", "").strip())


def read_case4_discrete_angles(
    workbook: str | Path,
    *,
    beam: str = "120 kVp",
    thickness_cm: float = 1.0,
    sheet: str = "Mean",
) -> dict[str, np.ndarray]:
    """The second test's per-projection-angle reference values, 36 per VOI.

    TG-195 runs the second test two ways (§3.4): (a) 36 discrete projections at
    10 degree increments, and (b) a source angle sampled randomly per photon from
    the continuous 360 degree distribution.  The published tables give both, and
    **they are not the same experiment**, so a 36-projection simulation must be
    compared against the discrete column.

    The two differ by 1.0-1.6% at the perimeter VOI in both beams, which is what a
    36-point quadrature of a sharply peaked angular function does -- the perimeter
    dose swings by a factor of 225 across the rotation.
    """
    import openpyxl

    key = (beam, float(thickness_cm))
    if key not in _CASE4_COLUMNS:
        raise KeyError(f"no Case 4 column for {key}; have {sorted(_CASE4_COLUMNS)}")
    col = _CASE4_COLUMNS[key]
    ws = openpyxl.load_workbook(Path(workbook), data_only=True)[sheet]

    out: dict[str, np.ndarray] = {}
    # The discrete block repeats the 'Center'/'Perimeter' labels after the
    # continuous one; the angle column identifies the rows that follow.
    seen: dict[str, int] = {}
    for r in range(1, ws.max_row + 1):
        label = ws.cell(row=r, column=2).value
        if not isinstance(label, str) or label.strip() not in ("Center", "Perimeter"):
            continue
        name = label.strip().lower()
        seen[name] = seen.get(name, 0) + 1
        if seen[name] < 2:  # first occurrence is the continuous case
            continue
        values = []
        for rr in range(r + 1, min(r + 45, ws.max_row + 1)):
            angle = ws.cell(row=rr, column=2).value
            value = ws.cell(row=rr, column=col).value
            if isinstance(angle, (int, float)) and isinstance(value, (int, float)):
                values.append(float(value))
            elif values:
                break
        if len(values) == len(CASE4_DISCRETE_PROJECTION_ANGLES_DEG):
            out[name] = np.asarray(values)
    return out


def read_case4_reference(
    workbook: str | Path,
    *,
    beam: str = "56.4 keV",
    thickness_cm: float = 1.0,
    angular: str = "continuous",
) -> dict[str, dict[str, float]]:
    """Reference values from ``09-Case 4 Results.xlsx``, in eV per photon.

    Takes the mean and the min-max range from AAPM's own ``Mean`` and ``Range``
    summary sheets rather than re-deriving them from the per-code sheets, whose
    layouts differ; the per-code values are attached as detail where they can be
    read.

    The **inter-code spread is the tolerance**.  Four independent packages agree
    with each other to 0.1-0.9% on the first test; agreeing with their mean any
    more tightly than they agree with each other would be luck, not validation.
    """
    import openpyxl

    key = (beam, float(thickness_cm))
    if key not in _CASE4_COLUMNS:
        raise KeyError(f"no Case 4 column for {key}; have {sorted(_CASE4_COLUMNS)}")
    col = _CASE4_COLUMNS[key]

    wb = openpyxl.load_workbook(Path(workbook), data_only=True)
    means = _read_case4_sheet(wb["Mean"], col)
    ranges = _read_case4_sheet(wb["Range"], col, numeric=False)
    per_code = {
        code: _read_case4_sheet(wb[code], col) for code in _CODES if code in wb.sheetnames
    }

    if angular not in ("continuous", "discrete"):
        raise ValueError("angular must be 'continuous' or 'discrete'")

    # Per-code discrete means, one sheet per package. These matter more than the
    # four-code average: the packages disagree with each other by 1.0% at the
    # centre VOI, so a single "deviation from the mean" hides the fact that the
    # reference is a spread rather than a value.
    per_code_discrete: dict[str, dict[str, float]] = {}
    if angular == "discrete":
        for code in _CODES:
            if code not in wb.sheetnames:
                continue
            angles = read_case4_discrete_angles(
                workbook, beam=beam, thickness_cm=thickness_cm, sheet=code
            )
            if angles:
                per_code_discrete[code] = {k: float(np.mean(v)) for k, v in angles.items()}
    discrete = (
        read_case4_discrete_angles(workbook, beam=beam, thickness_cm=thickness_cm)
        if angular == "discrete"
        else {}
    )

    # Which packages actually submitted the continuous sub-case. In the
    # distributed data only Geant4 and MCNP did, so the 'Mean' sheet's continuous
    # row is a TWO-code average while its discrete rows are four-code averages --
    # they are not two calculations of the same population.
    continuous_submitters = sorted(c for c, d in per_code.items() if d)

    out: dict[str, dict[str, float]] = {}
    for voi, mean in means.items():
        lo, hi = _parse_range(ranges[voi]) if voi in ranges else (mean, mean)
        rec: dict[str, float] = {
            "mean": float(mean),
            "min": lo,
            "max": hi,
            "inter_code_spread_fraction": (hi - lo) / mean if mean else float("nan"),
            "angular_case": "continuous",
            "continuous_submitted_by": continuous_submitters,
            "n_codes_continuous": len(continuous_submitters),
        }
        if voi in discrete:
            values = {c: d[voi] for c, d in per_code_discrete.items() if voi in d}
            rec["mean_continuous"] = float(mean)
            rec["mean"] = float(np.mean(discrete[voi]))
            rec["angular_case"] = "discrete (36-angle average)"
            rec["discrete_over_continuous"] = rec["mean"] / mean - 1.0
            rec["discrete_min"] = float(np.min(discrete[voi]))
            rec["discrete_max"] = float(np.max(discrete[voi]))
            if values:
                v = np.array(list(values.values()))
                rec["per_code"] = values
                rec["n_codes_discrete"] = len(values)
                rec["min"] = float(v.min())
                rec["max"] = float(v.max())
                rec["inter_code_spread_fraction"] = float((v.max() - v.min()) / v.mean())
                rec["inter_code_range_is_for_continuous_case"] = False
            else:  # pragma: no cover - only if the per-code sheets are missing
                rec["inter_code_range_is_for_continuous_case"] = True
        for code, d in per_code.items():
            if voi in d:
                rec[f"continuous_code_{code}"] = d[voi]
        out[voi] = rec
    return out


def _read_case4_sheet(ws, col: int, *, numeric: bool = True) -> dict[str, float]:
    """Pull the first-test segments and second-test VOIs out of one code's sheet.

    Located by label rather than by row number: the four sheets are not identical
    in layout, and a hard-coded row is the kind of thing that silently reads the
    wrong material.
    """
    labels = {}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=2).value
        if v is not None:
            labels.setdefault(str(v).strip(), r)

    def take(value):
        if numeric:
            return float(value) if isinstance(value, (int, float)) else None
        return value if isinstance(value, str) and "-" in value else None

    out: dict[str, float] = {}
    first = labels.get("Cylindrical Segment")
    if first is not None:
        for i in range(1, 5):
            # The segment number sits in column B, its value in `col`.
            for r in range(first, min(first + 12, ws.max_row + 1)):
                if ws.cell(row=r, column=2).value == i:
                    v = take(ws.cell(row=r, column=col).value)
                    if v is not None:
                        out[f"segment_{i}"] = v
                    break

    # "Center"/"Perimeter" appear twice: the continuous angular distribution
    # first, then the 36-angle discrete one. Take the first -- that is the case a
    # many-projection run should be compared against.
    seen: dict[str, list] = {"Center": [], "Perimeter": []}
    for r in range(1, ws.max_row + 1):
        v = ws.cell(row=r, column=2).value
        if isinstance(v, str) and v.strip() in seen:
            got = take(ws.cell(row=r, column=col).value)
            if got is not None:
                seen[v.strip()].append(got)
    if seen["Center"]:
        out["center"] = seen["Center"][0]
    if seen["Perimeter"]:
        out["perimeter"] = seen["Perimeter"][0]
    return out
